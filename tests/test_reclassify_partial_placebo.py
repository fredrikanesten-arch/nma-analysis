import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import reclassify_partial_placebo as module
from openpyxl import Workbook


class EnsureLocalFileTests(unittest.TestCase):
    def test_resolves_local_path(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            tmpdir = Path(temp_dir)
            local_file = tmpdir / "sample.txt"
            local_file.write_text("ok", encoding="utf-8")

            resolved = module.ensure_local_file(str(local_file), "unused.xlsx", tmpdir)

            self.assertEqual(resolved, local_file.resolve())

    @patch("reclassify_partial_placebo.urllib.request.urlretrieve")
    def test_downloads_default_remote_file(self, mock_urlretrieve):
        def fake_urlretrieve(url, destination):
            Path(destination).write_text("downloaded", encoding="utf-8")
            return str(destination), None

        mock_urlretrieve.side_effect = fake_urlretrieve
        with tempfile.TemporaryDirectory() as temp_dir:
            tmpdir = Path(temp_dir)

            resolved = module.ensure_local_file(None, "mmc5_fixed.xlsx", tmpdir)

            self.assertEqual(resolved, tmpdir / "mmc5_fixed.xlsx")
            mock_urlretrieve.assert_called_once_with(
                f"{module.DEFAULT_BASE_URL}/mmc5_fixed.xlsx",
                tmpdir / "mmc5_fixed.xlsx",
            )
            self.assertTrue(resolved.exists())

    def test_rejects_unapproved_remote_host(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            tmpdir = Path(temp_dir)
            with self.assertRaises(module.CliError):
                module.ensure_local_file("https://example.com/mmc5_fixed.xlsx", "mmc5_fixed.xlsx", tmpdir)


class SheetResolutionTests(unittest.TestCase):
    def test_resolve_mmc5_sheet_returns_exact_match(self):
        workbook = Workbook()
        workbook.active.title = "MS SMD bias-adj"

        resolved = module.resolve_mmc5_sheet(workbook, "MS SMD bias-adj")

        self.assertEqual(resolved, "MS SMD bias-adj")

    def test_resolve_mmc5_sheet_rejects_missing_sheet(self):
        workbook = Workbook()
        workbook.active.title = "MS SMD bias-adj"

        with self.assertRaises(module.CliError):
            module.resolve_mmc5_sheet(workbook, "MD SMD bias-adj")

    def test_infer_mmc3_sheet_supports_ms_and_ls(self):
        self.assertEqual(
            module.infer_mmc3_sheet("MS SMD bias-adj"),
            "MS depression-included studies",
        )
        self.assertEqual(
            module.infer_mmc3_sheet("LS SMD bias-adj"),
            "LS depression -included studies",
        )

    def test_infer_mmc3_sheet_rejects_unknown_prefix(self):
        with self.assertRaises(module.CliError):
            module.infer_mmc3_sheet("MD SMD bias-adj")


class RecodingTests(unittest.TestCase):
    def test_should_flag_requires_placebo_nonpharma_and_blinding_issue(self):
        flagged = module.StudyRecord(
            study_id="flagged",
            arms={1: "Pill placebo", 2: "Bright light therapy", 3: None, 4: None, 5: None},
            performance_bias="Unclear risk",
            detection_bias="Low risk",
        )
        not_flagged = module.StudyRecord(
            study_id="not_flagged",
            arms={1: "Pill placebo", 2: "Fluoxetine", 3: None, 4: None, 5: None},
            performance_bias="Low risk",
            detection_bias="Low risk",
        )

        self.assertTrue(module.should_flag(flagged))
        self.assertFalse(module.should_flag(not_flagged))

    def test_recode_sheet_updates_only_flagged_rows_and_tracks_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MS SMD bias-adj"
        headers = ["na[]", "t[,1]", "t[,2]", "t[,3]", "t[,4]", "t[,5]", "studyid"]
        for column, value in enumerate(headers, start=1):
            sheet.cell(1, column).value = value
        sheet.cell(2, 1).value = 2
        sheet.cell(2, 2).value = 1
        sheet.cell(2, 3).value = 42
        sheet.cell(2, 7).value = "S1"
        sheet.cell(3, 1).value = 2
        sheet.cell(3, 2).value = 1
        sheet.cell(3, 3).value = 42
        sheet.cell(3, 7).value = "S2"
        sheet.cell(4, 1).value = 3
        sheet.cell(4, 2).value = 1
        sheet.cell(4, 3).value = 1
        sheet.cell(4, 4).value = 42
        sheet.cell(4, 7).value = "S3"

        studies = {
            "S1": module.StudyRecord(
                study_id="S1",
                arms={1: "Pill placebo", 2: "Bright light therapy", 3: None, 4: None, 5: None},
                performance_bias="High risk",
                detection_bias="Low risk",
            ),
            "S2": module.StudyRecord(
                study_id="S2",
                arms={1: "Pill placebo", 2: "Fluoxetine", 3: None, 4: None, 5: None},
                performance_bias="Low risk",
                detection_bias="Low risk",
            ),
            "S3": module.StudyRecord(
                study_id="S3",
                arms={1: "Pill placebo", 2: "Pill placebo", 3: "Bright light therapy", 4: None, 5: None},
                performance_bias="High risk",
                detection_bias="Low risk",
            ),
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "mmc5_fixed.xlsx"
            workbook.save(workbook_path)

            recoded_path, flagged, _ = module.recode_sheet(
                workbook,
                workbook_path,
                studies,
                "MS SMD bias-adj",
                "MS depression-included studies",
            )

            recoded = Workbook()
            self.assertTrue(recoded_path.exists())
            recoded = module.load_workbook(recoded_path)
            recoded_sheet = recoded["MS SMD bias-adj"]
            self.assertEqual(recoded_sheet.cell(2, 2).value, module.PARTIAL_PLACEBO_CODE)
            self.assertEqual(recoded_sheet.cell(3, 2).value, module.PLACEBO_CODE)
            self.assertEqual(recoded_sheet.cell(4, 2).value, module.PARTIAL_PLACEBO_CODE)
            self.assertEqual(recoded_sheet.cell(4, 3).value, module.PARTIAL_PLACEBO_CODE)
            self.assertEqual([study.study_id for study in flagged], ["S1", "S3"])
            self.assertEqual(flagged[1].replaced_treat_columns, ["t[,1]", "t[,2]"])


if __name__ == "__main__":
    unittest.main()
