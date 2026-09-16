#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import shutil
import sys
import tempfile
import urllib.request
from urllib.parse import urlparse
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

DEFAULT_BASE_URL = "https://raw.githubusercontent.com/fredrikanesten-arch/placebo-reclassification/main"
DEFAULT_SHEET = "MS SMD bias-adj"
MMC3_SHEET_ALIASES = {
    "LS depression -included studies": "LS depression -included studies",
    "LS depression-included studies": "LS depression -included studies",
    "MS depression-included studies": "MS depression-included studies",
}
PLACEBO_CODE = 1
PARTIAL_PLACEBO_CODE = 100
PARTIAL_PLACEBO_NAME = "Partial placebo"
PLACEBO_CLASS_CODE = 1
PLACEBO_CLASS_NAME = "Placebo"
ARM_COLUMNS = [f"Arm {index} intervention" for index in range(1, 6)]
CONTROL_ARMS = {"pill placebo", "attention placebo", "no treatment", "waitlist", "tau"}
NONPHARMA_KEYWORDS = (
    "acupuncture",
    "attentional bias",
    "behavioral",
    "behavioural",
    "bibliotherapy",
    "bright light",
    "cbt",
    "cognitive bias",
    "computerised",
    "computerized",
    "counselling",
    "counseling",
    "dbt",
    "dialectical",
    "exercise",
    "interpersonal counselling",
    "interpersonal psychotherapy",
    "light therapy",
    "mbct",
    "meditation",
    "mindfulness",
    "music therapy",
    "peer support",
    "positive psychological",
    "problem solving",
    "psychoeducational",
    "psychodynamic",
    "psychotherapy",
    "relaxation",
    "self-help",
    "supportive",
    "therapy",
    "website",
    "yoga",
)
@dataclass
class StudyRecord:
    study_id: str
    arms: dict[int, str | None]
    performance_bias: str
    detection_bias: str


@dataclass
class FlaggedStudy:
    sheet_name: str
    block_index: int
    worksheet_row: int
    study_id: str
    matched_sheet: str
    replaced_treat_columns: list[str]
    performance_bias: str
    detection_bias: str
    arms: list[str]


class CliError(RuntimeError):
    pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Flag placebo-coded non-pharmacological studies with blinding issues "
            "and recode treatment 1 (Pill placebo) to 100 (Partial placebo)."
        )
    )
    parser.add_argument("--mmc5", help="Path or URL to mmc5_fixed.xlsx")
    parser.add_argument("--mmc3", help="Path or URL to mmc3_included_studies.xlsx")
    parser.add_argument("--lookup", help="Path or URL to trt_to_class_ms.csv")
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=(
            "Target mmc5 sheet."
        ),
    )
    parser.add_argument(
        "--mmc3-sheet",
        help="Optional override for the mmc3 sheet name; defaults to the sheet inferred from the resolved mmc5 sheet.",
    )
    parser.add_argument(
        "--output-dir",
        default="partial_placebo_outputs",
        help="Directory for the flagged-study report, recoded workbook, and updated lookup CSV.",
    )
    return parser.parse_args()


def ensure_local_file(source: str | None, default_name: str, tmpdir: Path) -> Path:
    if source is None:
        source = f"{DEFAULT_BASE_URL}/{default_name}"
    if source.startswith(("http://", "https://")):
        parsed = urlparse(source)
        expected_prefix = f"/fredrikanesten-arch/placebo-reclassification/main/{default_name}"
        if parsed.netloc != "raw.githubusercontent.com" or not parsed.path.endswith(expected_prefix):
            raise CliError(
                "Remote inputs must come from raw.githubusercontent.com/fredrikanesten-arch/"
                f"placebo-reclassification/main/{default_name} or be supplied as local files."
            )
        destination = tmpdir / default_name
        urllib.request.urlretrieve(source, destination)
        return destination
    return Path(source).expanduser().resolve()


def load_study_sheet(path: Path, sheet_name: str) -> dict[str, StudyRecord]:
    workbook = load_workbook(path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise CliError(f"Sheet '{sheet_name}' was not found in {path}.")
    worksheet = workbook[sheet_name]
    headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
    column_map = {header: index + 1 for index, header in enumerate(headers) if header}
    required_columns = [
        "Study ID",
        "Blinding of participants and personnel (performance bias)",
        "Blinding of outcome assessment (detection bias)",
    ]
    missing_columns = [name for name in required_columns if name not in column_map]
    if missing_columns:
        raise CliError(
            f"Sheet '{sheet_name}' in {path} is missing required columns: {', '.join(missing_columns)}"
        )
    study_id_col = column_map["Study ID"]
    perf_col = column_map["Blinding of participants and personnel (performance bias)"]
    det_col = column_map["Blinding of outcome assessment (detection bias)"]
    arm_cols = [column_map[name] for name in ARM_COLUMNS if name in column_map]
    studies: dict[str, StudyRecord] = {}
    for row in range(2, worksheet.max_row + 1):
        study_id = worksheet.cell(row, study_id_col).value
        if study_id is None:
            continue
        arms = {
            index: (
                None
                if worksheet.cell(row, column).value is None
                or normalize(str(worksheet.cell(row, column).value)) == "na"
                else str(worksheet.cell(row, column).value).strip()
            )
            for index, column in enumerate(arm_cols, start=1)
        }
        studies[str(study_id).strip()] = StudyRecord(
            study_id=str(study_id).strip(),
            arms=arms,
            performance_bias=str(worksheet.cell(row, perf_col).value or "").strip(),
            detection_bias=str(worksheet.cell(row, det_col).value or "").strip(),
        )
    return studies


def resolve_mmc5_sheet(workbook, requested_sheet: str) -> str:
    if requested_sheet in workbook.sheetnames:
        return requested_sheet
    raise CliError(
        f"Sheet '{requested_sheet}' was not found in mmc5 workbook. "
        f"Available sheets include: {', '.join(workbook.sheetnames)}"
    )


def infer_mmc3_sheet(mmc5_sheet_name: str) -> str:
    prefix = mmc5_sheet_name.split(" ", 1)[0]
    if prefix == "MS":
        return MMC3_SHEET_ALIASES["MS depression-included studies"]
    if prefix == "LS":
        return MMC3_SHEET_ALIASES["LS depression -included studies"]
    raise CliError(f"Unable to infer the mmc3 sheet for mmc5 sheet '{mmc5_sheet_name}'.")


def find_block_headers(worksheet) -> list[int]:
    return [row for row in range(1, worksheet.max_row + 1) if worksheet.cell(row, 1).value == "na[]"]


def build_column_map(worksheet, header_row: int) -> dict[str, int]:
    return {
        value: column
        for column in range(1, worksheet.max_column + 1)
        if (value := worksheet.cell(header_row, column).value) is not None
    }


def row_is_blank(worksheet, row: int) -> bool:
    return all(worksheet.cell(row, column).value is None for column in range(1, worksheet.max_column + 1))


def iter_non_responder_blocks(worksheet) -> Iterable[tuple[int, int, int, dict[str, int]]]:
    header_rows = find_block_headers(worksheet)
    block_index = 0
    for index, header_row in enumerate(header_rows):
        column_map = build_column_map(worksheet, header_row)
        if any(str(header).startswith("r[") for header in column_map):
            continue
        next_header = header_rows[index + 1] if index + 1 < len(header_rows) else worksheet.max_row + 1
        start_row = header_row + 1
        end_row = next_header - 1
        while end_row >= start_row and row_is_blank(worksheet, end_row):
            end_row -= 1
        block_index += 1
        yield block_index, start_row, end_row, column_map


def normalize(text: str) -> str:
    return " ".join(text.lower().split())


def split_arm_components(arm: str) -> list[str]:
    return [component.strip() for component in arm.split("+") if component.strip()]


def is_control_component(component: str) -> bool:
    return normalize(component) in CONTROL_ARMS


def is_nonpharmacological_component(component: str) -> bool:
    lowered = normalize(component)
    if lowered in CONTROL_ARMS:
        return False
    return any(keyword in lowered for keyword in NONPHARMA_KEYWORDS)


def study_has_nonpharmacological_component(study: StudyRecord) -> bool:
    for arm in study.arms.values():
        if arm is None:
            continue
        components = split_arm_components(arm)
        if any(is_nonpharmacological_component(component) for component in components):
            return True
    return False


def study_arms_list(study: StudyRecord) -> list[str]:
    return [arm for _, arm in sorted(study.arms.items()) if arm is not None]


def should_flag(study: StudyRecord) -> bool:
    normalized_arms = [normalize(arm) for arm in study_arms_list(study)]
    has_pill_placebo = "pill placebo" in normalized_arms
    has_nonpharma = study_has_nonpharmacological_component(study)
    has_blinding_issue = (
        normalize(study.performance_bias) != "low risk"
        or normalize(study.detection_bias) != "low risk"
    )
    return has_pill_placebo and has_nonpharma and has_blinding_issue


def pill_placebo_arm_count(study: StudyRecord) -> int:
    return sum(1 for arm in study.arms.values() if arm is not None and normalize(arm) == "pill placebo")


def recode_sheet(
    workbook,
    mmc5_path: Path,
    mmc3_studies: dict[str, StudyRecord],
    requested_sheet: str,
    mmc3_sheet_name: str,
) -> tuple[Path, list[FlaggedStudy], str]:
    sheet_name = resolve_mmc5_sheet(workbook, requested_sheet)
    worksheet = workbook[sheet_name]
    flagged: list[FlaggedStudy] = []
    for block_index, start_row, end_row, column_map in iter_non_responder_blocks(worksheet):
        treat_column_pairs = [
            (idx, column_map[name])
            for idx, name in enumerate([f"t[,{idx}]" for idx in range(1, 6)], start=1)
            if name in column_map
        ]
        treat_columns = [column for _, column in treat_column_pairs]
        study_id_column = column_map.get("studyid")
        if not treat_columns or study_id_column is None:
            continue
        for row in range(start_row, end_row + 1):
            row_codes = [worksheet.cell(row, column).value for column in treat_columns]
            if PLACEBO_CODE not in row_codes:
                continue
            study_id = worksheet.cell(row, study_id_column).value
            if study_id is None:
                continue
            study_key = str(study_id).strip()
            study = mmc3_studies.get(study_key)
            if study is None or not should_flag(study):
                continue
            placebo_positions_in_row = [position for position, column in treat_column_pairs if worksheet.cell(row, column).value == PLACEBO_CODE]
            expected_placebo_count = pill_placebo_arm_count(study)
            if len(placebo_positions_in_row) != expected_placebo_count:
                raise CliError(
                    f"Study '{study.study_id}' on sheet '{sheet_name}' row {row} has "
                    f"{len(placebo_positions_in_row)} treatment entries coded as {PLACEBO_CODE}, "
                    f"but mmc3 shows {expected_placebo_count} Pill placebo arm(s)."
                )
            changed = False
            changed_columns: list[str] = []
            for position, column in treat_column_pairs:
                if worksheet.cell(row, column).value == PLACEBO_CODE:
                    worksheet.cell(row, column).value = PARTIAL_PLACEBO_CODE
                    changed = True
                    changed_columns.append(f"t[,{position}]")
            if not changed:
                continue
            flagged.append(
                FlaggedStudy(
                    sheet_name=sheet_name,
                    block_index=block_index,
                    worksheet_row=row,
                    study_id=study.study_id,
                    matched_sheet=mmc3_sheet_name,
                    replaced_treat_columns=changed_columns,
                    performance_bias=study.performance_bias,
                    detection_bias=study.detection_bias,
                    arms=study_arms_list(study),
                )
            )
    recoded_path = mmc5_path.with_name(f"{mmc5_path.stem}_partial_placebo{mmc5_path.suffix}")
    workbook.save(recoded_path)
    return recoded_path, flagged, sheet_name


def write_flag_report(path: Path, flagged: list[FlaggedStudy]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "sheet_name",
                "block_index",
                "worksheet_row",
                "study_id",
                "matched_sheet",
                "replaced_treat_columns",
                "performance_bias",
                "detection_bias",
                "arms",
                "original_treat_code",
                "replacement_treat_code",
            ],
            quoting=csv.QUOTE_ALL,
        )
        writer.writeheader()
        for study in flagged:
            writer.writerow(
                {
                    "sheet_name": study.sheet_name,
                    "block_index": study.block_index,
                    "worksheet_row": study.worksheet_row,
                    "study_id": study.study_id,
                    "matched_sheet": study.matched_sheet,
                    "replaced_treat_columns": " | ".join(study.replaced_treat_columns),
                    "performance_bias": study.performance_bias,
                    "detection_bias": study.detection_bias,
                    "arms": " | ".join(study.arms),
                    "original_treat_code": PLACEBO_CODE,
                    "replacement_treat_code": PARTIAL_PLACEBO_CODE,
                }
            )


def update_lookup(source_path: Path, destination_path: Path) -> None:
    with source_path.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle, delimiter=";")
        rows = list(reader)
        fieldnames = [name.lstrip("\ufeff") for name in (reader.fieldnames or [])]
    if not fieldnames:
        raise CliError(f"Unable to read header row from {source_path}.")
    if rows:
        rows = [{key.lstrip("\ufeff"): value for key, value in row.items()} for row in rows]
    has_partial_placebo = any(row.get("trtcode") == str(PARTIAL_PLACEBO_CODE) for row in rows)
    if not has_partial_placebo:
        partial_placebo_row = {field: "" for field in fieldnames}
        values = {
            "trtcode": str(PARTIAL_PLACEBO_CODE),
            "trt": PARTIAL_PLACEBO_NAME,
            "classcode": str(PLACEBO_CLASS_CODE),
            "class": PLACEBO_CLASS_NAME,
        }
        for field in fieldnames:
            if field in values:
                partial_placebo_row[field] = values[field]
        rows.append(partial_placebo_row)
    with destination_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def stage_outputs(output_dir: Path, recoded_workbook: Path, flag_report: Path, updated_lookup: Path) -> tuple[Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_out = output_dir / recoded_workbook.name
    report_out = output_dir / flag_report.name
    lookup_out = output_dir / updated_lookup.name
    shutil.copy2(recoded_workbook, workbook_out)
    shutil.copy2(flag_report, report_out)
    shutil.copy2(updated_lookup, lookup_out)
    return workbook_out, report_out, lookup_out


def main() -> int:
    args = parse_args()
    output_dir = Path(args.output_dir).expanduser().resolve()
    with tempfile.TemporaryDirectory(prefix="partial-placebo-") as temp_dir:
        tmpdir = Path(temp_dir)
        mmc5_path = ensure_local_file(args.mmc5, "mmc5_fixed.xlsx", tmpdir)
        mmc3_path = ensure_local_file(args.mmc3, "mmc3_included_studies.xlsx", tmpdir)
        lookup_path = ensure_local_file(args.lookup, "trt_to_class_ms.csv", tmpdir)
        mmc5_sheet_name = args.sheet or DEFAULT_SHEET
        mmc5_workbook = load_workbook(mmc5_path)
        resolved_sheet = resolve_mmc5_sheet(mmc5_workbook, mmc5_sheet_name)
        resolved_mmc3_sheet = args.mmc3_sheet or infer_mmc3_sheet(resolved_sheet)
        mmc3_studies = load_study_sheet(mmc3_path, resolved_mmc3_sheet)
        recoded_workbook, flagged, resolved_sheet = recode_sheet(
            mmc5_workbook,
            mmc5_path,
            mmc3_studies,
            resolved_sheet,
            resolved_mmc3_sheet,
        )
        report_path = tmpdir / f"flagged_partial_placebo_{resolved_sheet.replace(' ', '_')}.csv"
        write_flag_report(report_path, flagged)
        lookup_out_path = tmpdir / f"{lookup_path.stem}_partial_placebo{lookup_path.suffix}"
        update_lookup(lookup_path, lookup_out_path)
        workbook_out, report_out, lookup_out = stage_outputs(output_dir, recoded_workbook, report_path, lookup_out_path)
        print(f"Resolved mmc5 sheet: {resolved_sheet}")
        print(f"Matched mmc3 sheet: {resolved_mmc3_sheet}")
        print(f"Flagged studies: {len(flagged)}")
        for study in flagged:
            print(
                f"- {study.study_id} (block {study.block_index}, row {study.worksheet_row}, "
                f"performance bias: {study.performance_bias})"
            )
        print(f"Recoded workbook: {workbook_out}")
        print(f"Flag report: {report_out}")
        print(f"Updated lookup: {lookup_out}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except CliError as error:
        print(f"error: {error}", file=sys.stderr)
        raise SystemExit(1)
